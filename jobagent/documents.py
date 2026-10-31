"""Tailored resume (DOCX + PDF) and cover letter for a job. Only resume facts are used; bullets are kept
verbatim and re-ordered by relevance, and any generated sentence that introduces a new number is rejected."""
import html
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor, Inches

from . import config, db, llm, resumes
from .keywords import FUNCTIONS


def _safe(s):
    return re.sub(r"[^\w.-]+", "_", s)[:60]


def job_dir(job):
    d = config.DOCS_DIR / f"{_safe(job['company'])}_{_safe(job['title'])}_{job['id'][:6]}"
    d.mkdir(parents=True, exist_ok=True)
    return d


def _allowed_numbers(job, key):
    prof = config.profile()
    nums = resumes.resume_numbers(key) | set(re.findall(r"\d+(?:\.\d+)?", job.get("description") or ""))
    nums |= {str(prof["total_experience_years"]), str(prof["post_mba_experience_years"])}
    return nums


def clean_title(title, company):
    """'Assistant Brand Manager, Sports, Amazon, Sports GL' -> 'Assistant Brand Manager, Sports'."""
    out = []
    for part in (p.strip() for p in title.split(",")):
        if not part or company.lower() in part.lower() or part in out:
            break
        out.append(part)
    return ", ".join(out[:2]) or title


def _lower_first(s):
    return s[0].lower() + s[1:] if s and not s[:2].isupper() else s


# ------------------------------------------------------------------ resume

def tailored_summary(t, job):
    """Profile paragraph. Claude may re-focus it on the JD; it falls back to the original wording."""
    original = t["profile"]
    if not llm.available():
        return original, "template"
    prompt = (f"Rewrite this resume profile summary (max 3 sentences, max 70 words) so it speaks to the job below. "
              f"Keep every claim traceable to the resume. Do not add metrics.\n\nPROFILE:\n{original}\n\n"
              f"RESUME FACTS:\n{resumes.load(t['key'])['full_text'][:6000]}\n\nJOB: {job['title']} at {job['company']}\n"
              f"{(job.get('description') or '')[:5000]}")
    out = llm.generate(prompt, max_tokens=2000)
    if not out or llm.invented_numbers(out, _allowed_numbers(job, t["key"])):
        return original, "template"
    return out.strip(), "claude"


def _docx_resume(t, summary, path):
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Inches(0.5)
        s.left_margin = s.right_margin = Inches(0.6)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)
    style.paragraph_format.space_after = Pt(1)

    def heading(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(7)
        r = p.add_run(text.upper())
        r.bold = True
        r.font.size = Pt(10.5)
        r.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(t["name"])
    r.bold = True
    r.font.size = Pt(18)
    prof = config.profile()
    contact = " | ".join(x for x in (prof.get("linkedin_url"), prof["email"], prof["phone"], prof["current_city"]) if x)
    p = doc.add_paragraph(contact)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    heading("Profile")
    doc.add_paragraph(summary)
    heading("Professional Experience")
    for i, role in enumerate(t["experience"]):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        title = t["current_title"] if i == 0 else role["title"]
        r = p.add_run(f"{title}, {role['company']}")
        r.bold = True
        p.add_run(f"    {role['start']} – {role['end']}").italic = True
        for b in role["bullets"]:
            doc.add_paragraph(b, style="List Bullet")
    if t["projects"] or t["projects_intro"]:
        heading("Product Training & Projects")
        if t["projects_intro"]:
            doc.add_paragraph(t["projects_intro"])
        for pr in t["projects"]:
            doc.add_paragraph(pr, style="List Bullet")
    heading("Skills")
    for s in t["skills"]:
        p = doc.add_paragraph()
        p.add_run(f"{s['category']}: ").bold = True
        p.add_run(", ".join(s["items"]))
    heading("Education")
    for e in t["education"]:
        doc.add_paragraph(e)
    doc.save(path)


def _html_resume(t, summary):
    e = html.escape
    prof = config.profile()
    contact = " &nbsp;|&nbsp; ".join(
        x for x in (f'<a href="{e(prof["linkedin_url"])}">LinkedIn</a>' if prof.get("linkedin_url") else "",
                    f'<a href="{e(prof["portfolio_url"])}">Product Portfolio</a>' if prof.get("portfolio_url") and t["key"] == "A" else "",
                    e(prof["email"]), e(prof["phone"]), e(prof["current_city"])) if x)
    roles = []
    for i, r in enumerate(t["experience"]):
        title = t["current_title"] if i == 0 else r["title"]
        roles.append(f'<div class="role"><b>{e(title)}, {e(r["company"])}</b><span>{e(r["start"])} – {e(r["end"])}</span></div>'
                     "<ul>" + "".join(f"<li>{e(b)}</li>" for b in r["bullets"]) + "</ul>")
    projects = ""
    if t["projects"] or t["projects_intro"]:
        projects = ("<h2>Product Training &amp; Projects</h2>" + (f"<p>{e(t['projects_intro'])}</p>" if t["projects_intro"] else "")
                    + "<ul>" + "".join(f"<li>{e(p)}</li>" for p in t["projects"]) + "</ul>")
    skills = "".join(f"<p><b>{e(s['category'])}:</b> {e(', '.join(s['items']))}</p>" for s in t["skills"])
    edu = "".join(f"<p>{e(x)}</p>" for x in t["education"])
    return f"""<!doctype html><html><head><meta charset="utf-8"><style>
    @page {{ size: A4; margin: 12mm 13mm; }}
    body {{ font-family: Calibri, 'Segoe UI', Arial, sans-serif; font-size: 10pt; color: #111; line-height: 1.32; }}
    h1 {{ text-align: center; font-size: 19pt; margin: 0; }}
    .contact {{ text-align: center; font-size: 9.5pt; margin: 2px 0 6px; }}
    a {{ color: #1f3a5f; text-decoration: none; }}
    h2 {{ font-size: 10.5pt; color: #1f3a5f; text-transform: uppercase; border-bottom: 1px solid #1f3a5f;
         margin: 8px 0 3px; padding-bottom: 1px; letter-spacing: .3px; }}
    .role {{ display: flex; justify-content: space-between; margin-top: 4px; }}
    .role span {{ font-style: italic; white-space: nowrap; padding-left: 8px; }}
    ul {{ margin: 1px 0 2px 16px; padding: 0; }} li {{ margin: 0 0 1px; }}
    p {{ margin: 1px 0; }}
    </style></head><body>
    <h1>{e(t['name'])}</h1><div class="contact">{contact}</div>
    <h2>Profile</h2><p>{e(summary)}</p>
    <h2>Professional Experience</h2>{''.join(roles)}
    {projects}<h2>Skills</h2>{skills}<h2>Education</h2>{edu}
    </body></html>"""


def html_to_pdf(html_text, path):
    """Render HTML to PDF with the installed Chrome/Edge through Playwright. Returns True on success."""
    from playwright.sync_api import sync_playwright
    with sync_playwright() as p:
        browser = None
        for channel in ("chrome", "msedge", None):
            try:
                browser = p.chromium.launch(channel=channel) if channel else p.chromium.launch()
                break
            except Exception:
                continue
        if browser is None:
            return False
        page = browser.new_page()
        page.set_content(html_text, wait_until="load")
        page.pdf(path=str(path), format="A4", print_background=True, prefer_css_page_size=True)
        browser.close()
    return True


# ------------------------------------------------------------------ cover letter

def _template_cover_letter(job, key, contact_name=None):
    prof = config.profile()
    r = resumes.load(key)
    t = resumes.tailor(key, job)
    current = r["experience"][0]
    top = sorted(resumes.achievements(key, job, n=3), key=lambda a: not a["role"].startswith(current["title"]))
    company = job["company"]
    fn_label = FUNCTIONS.get(job.get("function") or "", {}).get("label", "this")
    greeting = f"Dear {contact_name}," if contact_name else f"Dear {company} Hiring Team,"
    city = job.get("city")
    lines = [greeting, ""]
    lines.append(f"I am applying for the {clean_title(job['title'], company)} role{f' in {city}' if city and city != 'Remote' else ''}. "
                 f"I work at {current['company'].split(' (')[0]} as {t['current_title']}, and my experience maps "
                 f"closely to what the role asks for.")
    lines.append("")
    # Bullets keep their own tense, so lead-ins must work with both "own" and "owned".
    leads = iter(["In my current role, I", "I also", "Alongside this, I"])
    for a in top:
        b = _lower_first(a["bullet"]).rstrip(".")
        if a["role"].startswith(current["title"]):
            lines.append(f"{next(leads)} {b}.")
        else:
            title, _, org = a["role"].partition(", ")
            lines.append(f"Earlier, as {title} at {org}, I {b}.")
    lines.append("")
    if t["matched_skills"]:
        from .keywords import SKILLS
        jd = f"{job['title']} {job.get('description') or ''}"
        freq = {s: sum(len(re.findall(p, jd, re.I)) for p in SKILLS[s]) for s in t["matched_skills"]}
        sk = [s.lower() if not s[:3].isupper() else s for s in sorted(freq, key=lambda s: -freq[s])[:4]]
        joined = sk[0] if len(sk) == 1 else ", ".join(sk[:-1]) + f" and {sk[-1]}"
        lines.append(f"The role puts weight on {joined}; my resume shows hands-on work in each.")
    if config.pitch()["extras"]:
        lines.append(config.pitch()["extras"])
    lines.append("")
    loc_line = ""
    if city and city not in ("Remote", prof["current_city"]) and city in prof["preferred_locations"] + prof.get("open_locations", []):
        loc_line = f"I am based in {prof['current_city']} and open to relocating to {city}. "
    lines.append(f"{loc_line}I would welcome a conversation about how I can contribute to {company}'s {fn_label.lower()} work.")
    lines += ["", "Regards,", prof["name"], f"{prof['phone']} | {prof['email']}"]
    if prof.get("linkedin_url"):
        lines.append(prof["linkedin_url"])
    return "\n".join(lines)


def cover_letter(job, key, contact_name=None):
    """Returns (text, generator, warnings)."""
    template = _template_cover_letter(job, key, contact_name)
    if not llm.available():
        return template, "template", []
    prof = config.profile()
    prompt = (f"Write a cover letter (180-250 words) for this job. Address it to "
              f"{contact_name or 'the ' + job['company'] + ' hiring team'}. Pick the 2-3 resume achievements most "
              f"relevant to the JD and state them accurately. Mention relocation only if the job city differs from "
              f"{prof['current_city']}. End with the sign-off block exactly as in the DRAFT.\n\n"
              f"RESUME:\n{resumes.load(key)['full_text']}\n\nCANDIDATE: {prof['name']}, {prof['total_experience_years']} "
              f"years total experience, based in {prof['current_city']}.\n\nJOB: {job['title']} at {job['company']} "
              f"({job.get('location')})\n{(job.get('description') or '')[:6000]}\n\nDRAFT (facts are correct; improve the writing):\n{template}")
    out = llm.generate(prompt)
    if not out:
        return template, "template", ["Claude unavailable; used template"]
    bad = llm.invented_numbers(out, _allowed_numbers(job, key))
    if bad:
        return template, "template", [f"Claude draft rejected: numbers not in resume or JD ({', '.join(bad)})"]
    return out, "claude", []


def _docx_letter(text, path):
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(1)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    for para in text.split("\n"):
        doc.add_paragraph(para)
    doc.save(path)


# ------------------------------------------------------------------ entry point

def prepare(job_id, resume_key=None, contact_name=None):
    job = db.one("SELECT * FROM jobs WHERE id = ?", (job_id,))
    key = resume_key or job["resume_key"] or "A"
    t = resumes.tailor(key, job)
    summary, gen = tailored_summary(t, job)
    out = job_dir(job)
    role = re.sub(r"[,(].*$", "", job["title"]).strip()          # "Senior PM, Associate Satisfaction" -> "Senior PM"
    base = f"{_safe(prof['name'])}_{_safe(job['company'])}_{_safe(role)[:35].strip('_')}"
    files, warnings = [], []

    docx_path = out / f"{base}_Resume.docx"
    _docx_resume(t, summary, docx_path)
    files.append(("resume_docx", docx_path, gen))
    pdf_path = out / f"{base}_Resume.pdf"
    try:
        if html_to_pdf(_html_resume(t, summary), pdf_path):
            files.append(("resume_pdf", pdf_path, gen))
        else:
            warnings.append("No Chrome/Edge found for PDF export; upload the DOCX or the original PDF.")
    except Exception as e:
        warnings.append(f"PDF export failed: {e}")

    letter, lgen, lwarn = cover_letter(job, key, contact_name)
    warnings += lwarn
    txt_path = out / f"{base}_Cover_Letter.txt"
    txt_path.write_text(letter, encoding="utf-8")
    files.append(("cover_letter_txt", txt_path, lgen))
    cl_docx = out / f"{base}_Cover_Letter.docx"
    _docx_letter(letter, cl_docx)
    files.append(("cover_letter_docx", cl_docx, lgen))

    db.execute("DELETE FROM documents WHERE job_id = ?", (job_id,))
    for kind, path, g in files:
        db.insert("documents", {"job_id": job_id, "kind": kind, "path": str(path), "generator": g, "created_at": db.now()})
    fields = {"resume_key": key}
    if job["status"] in ("new", "shortlisted"):
        fields["status"] = "docs_ready"
    db.update("jobs", "id", job_id, fields)
    db.log_event(job_id, "documents_prepared", {"resume": key, "generator": gen, "warnings": warnings})
    return {"resume_key": key, "files": [{"kind": k, "path": str(p), "name": Path(p).name, "generator": g} for k, p, g in files],
            "matched_skills": t["matched_skills"], "missing_skills": t["missing_skills"], "warnings": warnings,
            "cover_letter": letter}
